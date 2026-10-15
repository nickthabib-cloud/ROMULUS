"""Professional DOCX and portable plain-text report generation."""
from pathlib import Path
from datetime import datetime, timezone
from typing import Any
from ..intelligence.knowledge import TECHNIQUES, ATTACK_RELEASE, ATTACK_RELEASE_DATE, ATTACK_DOMAIN

def generate_report(state: dict[str, Any], output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True,exist_ok=True)
    state=_xml_safe(state)
    hid=state.get("hunt_id","UNKNOWN"); stamp=datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    text_path=output_dir/f"{hid}-{stamp}.txt"; docx_path=output_dir/f"{hid}-{stamp}.docx"
    content=_content(state)
    text_path.write_text(content,encoding="utf-8")
    try: _docx(state,docx_path)
    except ImportError as e: raise RuntimeError("python-docx is required to generate the Word deliverable") from e
    return [docx_path,text_path]

def _xml_safe(value):
    """Remove characters forbidden by XML 1.0 from telemetry-derived text."""
    if isinstance(value,str):
        return "".join(ch for ch in value if ch in "\t\n\r" or 0x20<=ord(ch)<=0xD7FF or 0xE000<=ord(ch)<=0xFFFD or 0x10000<=ord(ch)<=0x10FFFF)
    if isinstance(value,dict):
        return {_xml_safe(key):_xml_safe(item) for key,item in value.items()}
    if isinstance(value,list):
        return [_xml_safe(item) for item in value]
    if isinstance(value,tuple):
        return tuple(_xml_safe(item) for item in value)
    return value

def _join_text(values, separator=", "):
    """Format model-derived lists defensively; structured output can still contain nulls."""
    if values is None: return ""
    if isinstance(values,str): return values
    if not isinstance(values,(list,tuple,set)): return str(values)
    return separator.join(str(value) for value in values if value is not None and str(value).strip())

def _event_reference(event):
    """Keep report citations useful without repeating large query DSL per event."""
    provenance=event.get("provenance",{}) if isinstance(event.get("provenance"),dict) else {}
    parts=[str(event.get("source") or "unknown")]
    retrieval=provenance.get("search_id") or provenance.get("lead_id")
    if retrieval: parts.append(f"search={retrieval}")
    for key in ("index","document_id","session_id","pcap"):
        value=provenance.get(key)
        if value is not None: parts.append(f"{key}={value}")
    raw=event.get("raw",{}) if isinstance(event.get("raw"),dict) else {}
    frame=raw.get("frame.number")
    if frame is not None: parts.append(f"frame={frame}")
    return "; ".join(parts)

def _bounded_report_events(state, key, limit=175):
    rows=list(state.get(key,[])); cited=set()
    for finding in state.get("findings",[]): cited.update(finding.get("evidence_ids",[]) or [])
    for hypothesis in state.get("hypotheses",[]): cited.update(hypothesis.get("evidence_ids",[]) or [])
    chosen=[row for row in rows if row.get("evidence_id") in cited]
    chosen_ids={row.get("evidence_id") for row in chosen}
    remaining=[row for row in rows if row.get("evidence_id") not in chosen_ids]
    slots=max(0,limit-len(chosen))
    if slots and remaining:
        step=max(1,len(remaining)//slots)
        chosen.extend(remaining[::step][:slots])
    return sorted(chosen[:max(limit,len(chosen))],key=lambda row:str(row.get("observed_at") or ""))

def _analytic_lines(summary):
    if not isinstance(summary,dict): return []
    lines=[]; scope=summary.get("scope",{}); quality=summary.get("telemetry_quality",{})
    lines.append(f"- Scope: host={scope.get('host_events',0)}, network={scope.get('network_events',0)}, first={scope.get('first_observed')}, last={scope.get('last_observed')}.")
    lines.append(f"- Data quality: timestamps={quality.get('timestamp_coverage_pct')}%, host entities={quality.get('host_entity_coverage_pct')}%, network endpoints={quality.get('network_endpoint_coverage_pct')}%.")
    for observation in summary.get("behavioral_observations",[])[:12]:
        lines.append(f"- {observation.get('analytic')}: observations={observation.get('observations')}; sample evidence={_join_text(observation.get('sample_evidence_ids'))}. {observation.get('note','')}")
    for fanout in summary.get("network_profile",{}).get("high_fanout",[])[:8]:
        lines.append(f"- Fan-out candidate {fanout.get('source')}: {fanout.get('unique_destinations')} destinations / {fanout.get('unique_destination_ports')} ports; sample={_join_text(fanout.get('sample_evidence_ids'))}.")
    for periodic in summary.get("network_profile",{}).get("periodic_edge_candidates",[])[:8]:
        lines.append(f"- Periodicity candidate {periodic.get('source')} → {periodic.get('destination')}:{periodic.get('port')}: {periodic.get('observations')} observations, mean {periodic.get('mean_interval_seconds')}s, CV {periodic.get('interval_cv')}; sample={_join_text(periodic.get('sample_evidence_ids'))}.")
    return lines


def _intelligence_coverage_lines(state):
    """Summarize rule execution without dumping the per-rule coverage ledger."""
    coverage = state.get("coverage", {}) if isinstance(state.get("coverage"), dict) else {}
    status = state.get("behavior_catalog_status", {})
    if not isinstance(status, dict):
        status = {}
    packs = status.get("behavior_packs", {})
    if not isinstance(packs, dict):
        packs = {}
    lines = []
    if packs:
        total = int(packs.get("curated_total") or 0)
        executable = int(packs.get("executable_curated") or 0)
        reference_only = int(packs.get("reference_only_curated") or max(0, total - executable))
        active = sum(int(packs.get(key) or 0) for key in (
            "active_direct", "active_sequences", "active_aggregates", "active_baselines"
        ))
        observe = int(packs.get("observe_curated") or 0)
        lines.append(
            f"Curated intelligence: loaded={total}; {executable} executable "
            f"({active} active, {observe} observe only); {reference_only} reference only; "
            f"capability references={int(packs.get('capability_references') or 0)}; "
            f"searches executed={int(status.get('searches_executed') or 0)}."
        )
    for label, key in (
        ("Sequence analytics", "sequence_analytics"),
        ("Aggregate analytics", "aggregate_analytics"),
        ("Baseline analytics", "baseline_analytics"),
    ):
        block = coverage.get(key, {})
        if not isinstance(block, dict) or not block:
            continue
        by_rule = block.get("by_rule", {})
        if not isinstance(by_rule, dict):
            by_rule = {}
        state_counts = {}
        for row in by_rule.values():
            value = row.get("state", "unknown") if isinstance(row, dict) else "unknown"
            state_counts[value] = state_counts.get(value, 0) + 1
        evaluated = sum(count for value, count in state_counts.items() if value != "reference_only")
        matched = state_counts.get("matched", 0) + state_counts.get("matched_partial", 0)
        partial = state_counts.get("partial", 0) + state_counts.get("matched_partial", 0)
        no_match = state_counts.get("no_match", 0)
        warming = state_counts.get("warming", 0)
        lines.append(
            f"{label}: loaded={int(block.get('total') or 0)}; "
            f"executable={int(block.get('executable') or 0)}; evaluated={evaluated}; "
            f"matched={matched}; partial={partial}; no match={no_match}; warming={warming}; "
            f"reference only={int(block.get('reference_only') or 0)}; "
            f"events evaluated={int(block.get('evaluated_events') or 0)}; "
            f"event limit reached={bool(block.get('event_limit_reached'))}."
        )
    return lines

def _candidate_line(candidate):
    return (
        f"- {candidate.get('candidate_id')} {candidate.get('title')} "
        f"[{candidate.get('priority_tier')} / score {candidate.get('priority_score')} / {candidate.get('disposition')}]; "
        f"kind={candidate.get('kind')}; evidence={_join_text(candidate.get('evidence_ids')) or 'none'}; "
        f"ATT&CK={_join_text(candidate.get('attack_techniques')) or 'unmapped'}; "
        f"score basis={_join_text(candidate.get('score_reasons'), '; ') or 'not recorded'}; "
        f"benign context={_join_text(candidate.get('benign_explanations'), '; ') or 'none supplied'}; "
        f"follow-up={_join_text(candidate.get('follow_up'), '; ') or 'none supplied'}."
    )


def _supported_paths(state):
    paths=list((state.get("investigation_graph") or {}).get("paths",[]) or [])
    referenced={
        path_id for finding in state.get("findings",[]) or []
        for path_id in (finding.get("attack_path_ids") or [])
    }
    if referenced:
        paths=[path for path in paths if path.get("path_id") in referenced]
    return sorted(paths,key=lambda path:(str(path.get("finding_id") or ""),int(path.get("segment_index") or 0),str(path.get("path_id") or "")))


def _path_entities(step):
    values=[]
    for entity in step.get("entities",[]) or []:
        value=entity.get("value")
        if value: values.append(f"{entity.get('entity_type','entity')}={value}")
    return _join_text(values)


def _scheduler_lines(state):
    tasks=list(state.get("search_tasks",{}).values()) if isinstance(state.get("search_tasks"),dict) else []
    outcomes={item.get("task_id"):item for item in state.get("search_outcomes",[]) or []
        if isinstance(item,dict) and item.get("task_id")}
    lines=[]
    terminal_resolved={"completed"}
    ordered=sorted(tasks,key=lambda item:(
        1 if (outcomes.get(item.get("task_id"),{}).get("status") or item.get("status")) in terminal_resolved else 0,
        int(item.get("created_sequence",0)),str(item.get("task_id",""))))
    selected=ordered[:250]
    lines.append(f"Scheduler task appendix: shown={len(selected)}; total={len(tasks)}; omitted={max(0,len(tasks)-len(selected))}.")
    for task in selected:
        outcome=outcomes.get(task.get("task_id"),{})
        sampling=outcome.get("sampling") if isinstance(outcome.get("sampling"),dict) else {}
        retained=sampling.get("retained_by_stratum")
        matched_by=sampling.get("matched_by_stratum")
        retained_total=sum(value for value in retained if isinstance(value,int) and not isinstance(value,bool)) if isinstance(retained,list) else 0
        sampling_text=(f"sampling={sampling.get('strategy')}; strata={sampling.get('covered_strata')}/{sampling.get('requested_strata')}; "
            f"sampled_records={retained_total}; retained_by_stratum={_join_text(retained) or 'not recorded'}; "
            f"matched_by_stratum={_join_text(matched_by) or 'not recorded'}; "
            f"sampling_fallback={sampling.get('fallback_reason') or 'none'}; "
            if sampling else "sampling=not recorded; ")
        owners=[]
        for key,label in (("candidate_ids","candidates"),("follow_up_ids","follow-ups"),("hypothesis_ids","hypotheses")):
            if task.get(key): owners.append(f"{label}={_join_text(task.get(key))}")
        scope=(outcome.get("time_range") or task.get("time_range") or {}).get("scope","not recorded")
        lines.append(f"- {task.get('task_id')} {task.get('source')}/{task.get('source_family')} {task.get('phase')}; "
            f"origins={_join_text(task.get('origins') or [task.get('origin')])}; intent={task.get('intent')}; "
            f"owners={'; '.join(owners) or 'none'}; status={outcome.get('status') or task.get('status')}; "
            f"returned={outcome.get('returned_total')}; matched={outcome.get('matched_total')}; scope={scope}; "
            f"capped={bool(outcome.get('capped'))}; timed_out={bool(outcome.get('timed_out'))}; "
            f"shard_failures={outcome.get('shard_failures',0)}; incomplete={bool(outcome.get('incomplete'))}; "
            f"{sampling_text}error={outcome.get('error_kind') or outcome.get('reason') or 'none'}.")
    gaps=[item for item in state.get("scheduler_history",[]) or []
        if isinstance(item,dict) and "budget_gap" in str(item.get("event"))]
    gaps.extend(state.get("coverage",{}).get("scheduler_gaps",[]) or [])
    if gaps:
        lines.append("Unresolved scheduler/source-fairness gaps:")
        for gap in gaps[-40:]:
            task_ids=gap.get("budget_limited_task_ids") or gap.get("queued_task_ids") or []
            total=gap.get("budget_limited_task_count")
            total=int(total) if isinstance(total,int) else len(task_ids)
            lines.append(f"- phase={gap.get('phase')}; event={gap.get('event')}; reason={gap.get('reason')}; "
                f"unserved={_join_text(gap.get('unserved_families')) or 'none'}; "
                f"budget-limited={_join_text(task_ids) or 'none'}; total={total}; omitted={max(0,total-len(task_ids))}.")
    return lines


def _content(s):
    host_events=_bounded_report_events(s,"host"); network_events=_bounded_report_events(s,"network")
    lines=[f"ROMULUS INVESTIGATION REPORT — {s.get('status','partial').upper()}",f"Investigation: {s.get('hunt_id','')}",f"Generated: {datetime.now(timezone.utc).isoformat(timespec='seconds')}","", "EXECUTIVE SUMMARY",s.get("summary") or "Investigation summary unavailable.",
        f"Outcome: {s.get('outcome','not recorded')}",f"Termination reason: {s.get('termination_reason','not recorded')}",
        f"Elapsed: {s.get('duration_ms','not recorded')} ms",
        f"Tool calls: {s.get('tool_calls',0)} / {s.get('coverage',{}).get('tool_call_limit','not recorded')}; analyst rounds: {s.get('depth',0)} / {s.get('coverage',{}).get('depth_limit','not recorded')}",
        f"ATT&CK knowledge baseline: {ATTACK_DOMAIN} {ATTACK_RELEASE} ({ATTACK_RELEASE_DATE})","", "OBJECTIVE AND SCOPE",s.get("objective",""),"", "DATA SOURCES AND COVERAGE"]
    verbose_coverage = {
        "search_fields", "analytic_summary", "hypothesis_coverage", "recent_search_artifacts",
        "behavior_catalog", "search_matrix", "sequence_analytics", "aggregate_analytics",
        "baseline_analytics",
    }
    lines += [f"- {k}: {v}" for k,v in s.get("coverage",{}).items() if k not in verbose_coverage]
    lines += ["", "INTELLIGENCE EXECUTION"]
    lines += [f"- {line}" for line in _intelligence_coverage_lines(s)] or ["No intelligence execution coverage was recorded."]
    lines += ["", "REASONING AND INVESTIGATION STAGES"]
    for stage in s.get("stage_history",[]): lines.append(f"- {stage.get('timestamp')} {stage.get('stage')}: {stage.get('status')}; details={stage.get('details')}.")
    for stage in s.get("reasoning_history",[]):
        searches=[x.get("intent",x.get("id","")) for x in stage.get("selected_searches",[]) if isinstance(x,dict)]
        lines.append(f"- {stage.get('stage')}: {stage.get('basis')}; round={stage.get('round','n/a')}; next searches={_join_text(searches, '; ') or _join_text(stage.get('selected_leads')) or 'none'}; evidence considered={stage.get('evidence_count',0)}.")
        if stage.get("reasoning"): lines.append(f"  Analyst reasoning: {stage.get('reasoning')}")
        if stage.get("created_hypothesis_ids"): lines.append(f"  New follow-up hypotheses: {_join_text(stage.get('created_hypothesis_ids'))}")
    for run in s.get("tool_runs",[]): lines.append(f"- {run.get('reasoning_stage')}: {run.get('tool')} [{run.get('result_status')}], returned={run.get('result_count',0)}, matched total={run.get('source_total')}, window={run.get('time_range')}, duration={run.get('duration_ms')} ms, error={run.get('error') or 'none'}.")
    lines += ["", "HYPOTHESES"]
    for h in s.get("hypotheses",[]):
        t=TECHNIQUES.get(h.get("technique_id"),{}); lines.append(f"- {h.get('title')} ({h.get('technique_id') or 'unmapped'} {t.get('name','')}): {h.get('status')}; confidence {h.get('confidence')}. {h.get('rationale')}")
        lines.append(f"  Expected host evidence: {_join_text(h.get('expected_host'))}; network: {_join_text(h.get('expected_network'))}. Benign explanations: {_join_text(h.get('benign_explanations'))}.")
        assessment=h.get("model_assessment") or h.get("deterministic_assessment")
        if assessment: lines.append(f"  Evidence assessment: {assessment.get('status')}; rationale={assessment.get('rationale')}; cited events={_join_text(assessment.get('evidence_ids'))}.")
    lines += ["", "HYPOTHESIS AND ANALYTIC COVERAGE"]
    ledger=s.get("coverage",{}).get("hypothesis_coverage",{})
    for row in ledger.get("rows",[]):
        lines.append(f"- {row.get('hypothesis_id')} / {row.get('technique_id')}: priority={row.get('priority')}; analytics attempted={row.get('executed_analytics')}/{row.get('planned_analytics')}; resolved={row.get('resolved_analytics')}; sources={_join_text(row.get('sources')) or 'none'}; returned={row.get('returned')}; assessment={row.get('assessment')}; gaps={_join_text(row.get('telemetry_gaps')) or 'none'}.")
    lines += ["", "UNIFIED SEARCH SCHEDULER"]+(_scheduler_lines(s) or ["No scheduler tasks were recorded."])
    lines += ["", "DETERMINISTIC BEHAVIORAL ANALYTICS"]+(_analytic_lines(s.get("analytic_summary",{})) or ["No analytic summary was produced."])
    visible_candidates=[item for item in s.get("candidate_leads",[]) if item.get("customer_visible")]
    lines += ["", "PRIORITIZED CANDIDATE LEADS"]
    lines += [_candidate_line(item) for item in visible_candidates] or ["No trusted or corroborated high-priority catalog leads met the customer-report threshold."]
    lines += ["", "SIGNIFICANT FINDINGS"]
    if not s.get("findings"): lines.append("No significant findings recorded.")
    for f in s.get("findings",[]): lines += [f"- {f.get('finding_id')} {f.get('title')} [{f.get('status')} / {f.get('confidence')}]: {f.get('observation')}","  Analysis: "+str(f.get("interpretation") or ""),"  Evidence: "+_join_text(f.get("evidence_ids")),"  Alternatives: "+_join_text(f.get("alternatives"),"; "),"  Follow-up: "+_join_text(f.get("follow_up"),"; ")]
    paths=_supported_paths(s)
    lines += ["", "SUPPORTED ATTACK PATHS"]
    if not paths: lines.append("No evidence-supported multi-record path was derived from the cited findings.")
    for path in paths:
        lines.append(f"- {path.get('path_id')} {path.get('title')} — segment {path.get('segment_index',1)} of {path.get('segment_count',1)}; relationship confidence={path.get('confidence','not assessed')}.")
        for step in path.get("steps",[]):
            lines.append(f"  {step.get('sequence')}. {step.get('observed_at') or 'time unavailable'} {step.get('activity_type')}: {step.get('label')} Evidence={_join_text(step.get('evidence_ids'))}; entities={_path_entities(step) or 'none recorded'}.")
        for relationship in path.get("relationships",[]):
            lines.append(f"  Link {relationship.get('source_sequence')}→{relationship.get('target_sequence')}: {relationship.get('relationship')} [{relationship.get('basis')} / {relationship.get('confidence')}]; reasons={_join_text(relationship.get('reasons'), '; ') or 'not recorded'}.")
        for caveat in path.get("caveats",[]): lines.append(f"  Caveat: {caveat}")
        for gap in path.get("gaps",[]): lines.append(f"  Gap: {gap}")
    lines += ["", f"HOST ANALYSIS — {len(host_events)} representative/cited records of {len(s.get('host',[]))}"]+[f"- {e.get('evidence_id')} {e.get('observed_at')} {e.get('asset')}: {e.get('observation')} [{_event_reference(e)}]" for e in host_events]
    lines += ["", f"NETWORK ANALYSIS — {len(network_events)} representative/cited records of {len(s.get('network',[]))}"]+[f"- {e.get('evidence_id')} {e.get('observed_at')} {e.get('asset')}: {e.get('observation')} [{_event_reference(e)}]" for e in network_events]
    lines += ["", "CROSS-DOMAIN CORRELATION"]
    lines += [f"- Host {c.get('host_ref')} ↔ network {c.get('network_ref')}; reasons={_join_text(c.get('reasons'), '; ') or ('shared IP '+str(c.get('ip')))}; direction={c.get('direction','unresolved')}; score={c.get('score','not scored')}; time delta {c.get('delta_seconds')} seconds. Caveat: {c.get('caveat')}" for c in s.get("correlations",[])] or ["No corroborated cross-domain correlation was made."]
    lines += ["", "ATT&CK MAPPING"]+[f"- {h.get('technique_id')} {TECHNIQUES.get(h.get('technique_id'),{}).get('name','')}: {h.get('status')}" for h in s.get("hypotheses",[])]
    lines += ["", "SEARCH RESULT ARTIFACTS"]
    lines += [f"- {a.get('artifact_id')} {a.get('stage')} via {a.get('tool')}: status={a.get('result_status')}; returned={a.get('returned')}; matched={a.get('matched_total')}; scope={a.get('scope')}; error={a.get('error_reason') or 'none'}." for a in s.get("search_artifacts",[])[:250]] or ["No search artifacts were recorded."]
    lines += ["", "BOUNDED EVIDENCE TIMELINE"]
    timeline=sorted(host_events+network_events,key=lambda e:e.get("observed_at","") or "")
    lines += [f"- {e.get('observed_at')} {e.get('source')}: {e.get('observation')} ({e.get('evidence_id')})" for e in timeline]
    lines += ["", "SOURCE ERRORS AND LIMITATIONS"]
    errors=s.get("source_errors",[])
    lines += [f"- {e.get('tool')} at {e.get('stage')}: {e.get('error')}; {e.get('reason')} (HTTP {e.get('status_code')})" if e.get("status_code") else f"- {e.get('tool')} at {e.get('stage')}: {e.get('error')}; {e.get('reason') or 'no additional detail'}" for e in errors] or ["No source errors were recorded."]
    if s.get("error"): lines.append(s["error"])
    lines += ["This assessment covers only the evidence sources and time period recorded above. Searches and reports are bounded; lack of findings does not establish that the environment is secure.","", "CONCLUSIONS AND RECOMMENDED FOLLOW-UP",s.get("summary") or "", "", "TECHNICAL CANDIDATE WORKING LIST"]
    lines += [_candidate_line(item) for item in s.get("candidate_leads",[])] or ["No catalog candidates were recorded."]
    if s.get("retention_policy")=="compact" and (s.get("status")=="completed" or s.get("finalized_at")):
        retention_note=("This finalized investigation uses compact retention. After both reports are verified, per-hunt working audit records and uncited raw evidence are scheduled for removal; "
            "structured findings, supported paths, and their cited evidence remain in the compact investigation record.")
    else:
        retention_note="Exact compiled queries remain in tool_audit.jsonl and complete raw evidence remains in SQLite unless the hunt is explicitly finalized with compact retention."
    lines += ["", "TECHNICAL EVIDENCE APPENDIX", f"The human-readable evidence manifest is bounded to {len(host_events)} host and {len(network_events)} network records. Stable citations and search-result IDs appear above. {retention_note}"]
    return "\n".join(lines)

def _docx(s,path):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    navy="18324A"; blue="245B83"; pale_blue="EAF2F8"; pale_gray="F3F5F7"; white="FFFFFF"; amber="C77800"; green="2F6B4F"

    def shade(cell,fill):
        properties=cell._tc.get_or_add_tcPr(); element=properties.find(qn("w:shd"))
        if element is None:
            element=OxmlElement("w:shd"); properties.append(element)
        element.set(qn("w:fill"),fill)

    def set_cell_text(cell,text,bold=False,color=None,size=9):
        cell.text=""; paragraph=cell.paragraphs[0]; run=paragraph.add_run(str(text or "Not recorded"))
        run.bold=bold; run.font.size=Pt(size)
        if color: run.font.color.rgb=RGBColor.from_string(color)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def mark_header(row):
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    def format_table(table,header=True):
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for row_index,row in enumerate(table.rows):
            for cell in row.cells:
                properties=cell._tc.get_or_add_tcPr()
                borders=properties.find(qn("w:tcBorders"))
                if borders is None:
                    borders=OxmlElement("w:tcBorders"); properties.append(borders)
                for edge in ("top","left","bottom","right","insideH","insideV"):
                    border=borders.find(qn(f"w:{edge}"))
                    if border is None:
                        border=OxmlElement(f"w:{edge}"); borders.append(border)
                    border.set(qn("w:val"),"single"); border.set(qn("w:sz"),"4"); border.set(qn("w:color"),"D9D9D9")
                margins=properties.find(qn("w:tcMar"))
                if margins is None:
                    margins=OxmlElement("w:tcMar"); properties.append(margins)
                for side in ("top","left","bottom","right"):
                    margin=margins.find(qn(f"w:{side}"))
                    if margin is None:
                        margin=OxmlElement(f"w:{side}"); margins.append(margin)
                    margin.set(qn("w:w"),"90"); margin.set(qn("w:type"),"dxa")
                if header and row_index==0:
                    shade(cell,navy)
                elif row_index%2==0:
                    shade(cell,pale_gray)

    def heading(title):
        paragraph=doc.add_heading(title,level=1); paragraph.paragraph_format.space_before=Pt(13); paragraph.paragraph_format.space_after=Pt(5)
        return paragraph

    def para(value,style=None):
        paragraph=doc.add_paragraph(str(value or "Not recorded."),style=style)
        paragraph.paragraph_format.space_after=Pt(5)
        return paragraph

    def friendly(value):
        return str(value or "Not recorded").replace("_"," ").strip().capitalize()

    findings=list(s.get("findings",[]) or [])
    visible_candidates=[item for item in s.get("candidate_leads",[]) if item.get("customer_visible")]
    errors=list(s.get("source_errors",[]) or [])
    sampled_outcomes=[item for item in s.get("search_outcomes",[]) or []
        if isinstance(item,dict) and (item.get("capped") or
            isinstance(item.get("sampling"),dict) and item["sampling"].get("strategy") not in {None,"exhaustive"})]
    representative_strategies={"time_stratified","rank_stratified","time_stratified_reservoir","deterministic_reservoir"}
    representative_outcomes=[item for item in sampled_outcomes
        if item.get("sampling",{}).get("strategy") in representative_strategies
        and not item.get("sampling",{}).get("fallback_reason")]
    fallback_outcomes=[item for item in sampled_outcomes if item not in representative_outcomes]
    confidence_order={"not assessed":0,"low":1,"moderate":2,"high":3}
    overall_confidence=max((str(item.get("confidence") or "low").lower() for item in findings),key=lambda value:confidence_order.get(value,0),default="low")
    if s.get("status")!="completed" or errors:
        verdict="Inconclusive"
    elif findings or visible_candidates:
        verdict="Suspicious Activity Identified"
    else:
        verdict="No Significant Findings"
    tiers=[str(item.get("priority_tier") or "").lower() for item in visible_candidates]
    priority="Critical" if "critical" in tiers else "High" if "high" in tiers else "Moderate" if findings else "Informational"

    cited_ids={evidence for finding in findings for evidence in (finding.get("evidence_ids") or [])}
    assets=[]
    for finding in findings:
        for key in ("affected_assets","assets"):
            value=finding.get(key) or []
            if isinstance(value,str): value=[value]
            assets.extend(str(item) for item in value if item)
    for event in list(s.get("host",[]) or [])+list(s.get("network",[]) or []):
        if event.get("evidence_id") in cited_ids and event.get("asset"): assets.append(str(event["asset"]))
    assets=list(dict.fromkeys(assets))

    represented_titles={str(item.get("title") or "").strip().lower() for item in findings}
    action_items=findings+[item for item in visible_candidates if str(item.get("title") or "").strip().lower() not in represented_titles]
    actions=[]
    for item in action_items:
        values=item.get("follow_up") or []
        if isinstance(values,str): values=[values]
        actions.extend(str(value).strip() for value in values if str(value).strip())
    actions=list(dict.fromkeys(actions))[:6]
    if not actions:
        actions=["Review the hunt scope and retain the report with the investigation record."]

    attack_rows=[]; seen_techniques=set()
    technique_inputs=[]
    for finding in findings:
        technique_inputs.extend([finding.get("technique_id")]+list(finding.get("attack_techniques") or []))
    for candidate in visible_candidates: technique_inputs.extend(candidate.get("attack_techniques") or [])
    for hypothesis in s.get("hypotheses",[]) or []:
        if str(hypothesis.get("status") or "").lower() in {"supported","suspected","observed"}:
            technique_inputs.append(hypothesis.get("technique_id"))
    for technique_id in technique_inputs:
        if not technique_id or technique_id in seen_techniques or technique_id not in TECHNIQUES: continue
        seen_techniques.add(technique_id); technique=TECHNIQUES[technique_id]
        attack_rows.append((technique_id,technique.get("name") or "Not recorded",technique.get("tactic") or "Not recorded"))

    doc=Document(); section=doc.sections[0]
    section.top_margin=Inches(.6); section.bottom_margin=Inches(.6); section.left_margin=Inches(.72); section.right_margin=Inches(.72)
    header=section.header.paragraphs[0]; header.text="ROMULUS  |  DEFENSIVE THREAT HUNT"; header.style=doc.styles["Header"]
    header.runs[0].font.color.rgb=RGBColor(0,0,0); header.runs[0].font.size=Pt(8); header.runs[0].bold=True
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{s.get('hunt_id','UNKNOWN')}  •  UNMARKED  •  Page ").font.size=Pt(8)
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)

    styles=doc.styles; styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(9.5); styles["Normal"].font.color.rgb=RGBColor.from_string(navy)
    styles["Title"].font.name="Aptos Display"; styles["Title"].font.size=Pt(25); styles["Title"].font.bold=True; styles["Title"].font.color.rgb=RGBColor(0,0,0)
    styles["Heading 1"].font.name="Aptos Display"; styles["Heading 1"].font.size=Pt(15); styles["Heading 1"].font.bold=True; styles["Heading 1"].font.color.rgb=RGBColor(0,0,0)
    styles["Heading 2"].font.name="Aptos"; styles["Heading 2"].font.size=Pt(11); styles["Heading 2"].font.bold=True; styles["Heading 2"].font.color.rgb=RGBColor(0,0,0)

    title_properties=styles["Title"].element.get_or_add_pPr(); title_border=title_properties.find(qn("w:pBdr"))
    if title_border is not None: title_properties.remove(title_border)
    title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.LEFT; title.add_run("Threat Hunt Report")
    subtitle=doc.add_paragraph(); subtitle.paragraph_format.space_after=Pt(12)
    run=subtitle.add_run("ROMULUS autonomous investigation  •  "); run.bold=True; run.font.color.rgb=RGBColor.from_string(blue)
    subtitle.add_run(f"{s.get('hunt_id','UNKNOWN')}  •  {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")

    assessment=doc.add_paragraph(); assessment.paragraph_format.space_after=Pt(12)
    label=assessment.add_run("Assessment  "); label.bold=True
    value=assessment.add_run(verdict); value.bold=True; value.font.color.rgb=RGBColor.from_string(amber if verdict!="No Significant Findings" else green)
    assessment.add_run(f"    Confidence  {overall_confidence.title()}    Priority  {priority}")

    heading("Executive Summary"); para(s.get("summary") or "No investigation summary was recorded.")
    heading("Assessment at a Glance")
    glance=doc.add_table(rows=0,cols=2); glance.style="Light Shading Accent 1"; glance.alignment=WD_TABLE_ALIGNMENT.CENTER
    glance_items=(
        ("Objective",s.get("objective") or "Not recorded"),
        ("Hunt window",s.get("coverage",{}).get("time_scope") or "Not recorded"),
        ("Affected assets",", ".join(assets) if assets else "No affected assets identified"),
        ("Finding count",str(len(findings))),
    )
    for label,value in glance_items:
        cells=glance.add_row().cells; set_cell_text(cells[0],label,True,navy); set_cell_text(cells[1],value)
    format_table(glance,header=False)

    heading("What We Found")
    if not findings and not visible_candidates:
        para("No reportable suspicious activity met the finding threshold in the telemetry and time range examined.")
    finding_titles=set()
    for finding in findings:
        title_text=str(finding.get("title") or "Untitled finding"); finding_titles.add(title_text.lower())
        p=doc.add_heading(title_text,level=2)
        meta=doc.add_paragraph(f"{friendly(finding.get('status'))}  •  {friendly(finding.get('confidence'))} confidence")
        meta.runs[0].bold=True; meta.runs[0].font.color.rgb=RGBColor.from_string(blue)
        para(str(finding.get("observation") or "No observation was recorded."))
        if finding.get("interpretation"): para(str(finding["interpretation"]))
        alternatives=_join_text(finding.get("alternatives"),"; ")
        if alternatives: para("Possible benign explanation: "+alternatives)
    for candidate in visible_candidates:
        if str(candidate.get("title") or "").lower() in finding_titles: continue
        doc.add_heading(str(candidate.get("title") or "Priority lead"),level=2)
        para(f"Priority lead  •  {friendly(candidate.get('priority_tier'))} priority  •  {friendly(candidate.get('disposition'))}")
        reasons=_join_text(candidate.get("score_reasons"),"; ")
        if reasons: para("Why it matters: "+reasons)

    heading("Supported Attack Path")
    supported_paths=_supported_paths(s)
    if not supported_paths:
        para("No evidence-supported multi-record path was derived from the cited findings.")
    for supported_path in supported_paths:
        segment=f" — segment {supported_path.get('segment_index',1)} of {supported_path.get('segment_count',1)}" if int(supported_path.get("segment_count") or 1)>1 else ""
        doc.add_heading(str(supported_path.get("title") or "Evidence-supported activity")+segment,level=2)
        for step in supported_path.get("steps",[]):
            entities=_path_entities(step)
            detail=f"{step.get('sequence')}. {step.get('observed_at') or 'Time unavailable'} — {friendly(step.get('activity_type'))}: {step.get('label')}. Evidence: {_join_text(step.get('evidence_ids'))}."
            if entities: detail+=f" Entities: {entities}."
            para(detail)
        caveats=list(supported_path.get("caveats",[]) or [])+list(supported_path.get("gaps",[]) or [])
        if caveats: para("Limitations: "+_join_text(caveats," "))

    heading("Scope and Coverage")
    scope=doc.add_table(rows=1,cols=2); scope.style="Light Shading Accent 1"; scope.alignment=WD_TABLE_ALIGNMENT.CENTER
    set_cell_text(scope.rows[0].cells[0],"Telemetry",True,white); set_cell_text(scope.rows[0].cells[1],"Coverage result",True,white)
    shade(scope.rows[0].cells[0],navy); shade(scope.rows[0].cells[1],navy); mark_header(scope.rows[0])
    coverage=s.get("coverage",{}) or {}
    for label,key in (("Host telemetry","host"),("Network telemetry","network")):
        cells=scope.add_row().cells; set_cell_text(cells[0],label,True); set_cell_text(cells[1],coverage.get(key) or "Not recorded")
    limitations=[]
    if errors: limitations.append(f"{len(errors)} source error(s) limited the investigation.")
    if representative_outcomes:
        limitations.append(f"{len(representative_outcomes)} search result set(s) used a bounded representative sample; capped searches remain incomplete coverage.")
    if fallback_outcomes:
        limitations.append(f"{len(fallback_outcomes)} search result set(s) used incomplete, nonrepresentative bounded retrieval because complete representative sampling was unavailable.")
    if not limitations: limitations.append("No source errors or sampled-result limitations were recorded.")
    cells=scope.add_row().cells; set_cell_text(cells[0],"Collection limitations",True)
    set_cell_text(cells[1]," ".join(limitations)+" See the technical TXT report for details.")
    format_table(scope)
    doc.add_heading("Intelligence execution",level=2)
    intelligence_lines=_intelligence_coverage_lines(s)
    for line in intelligence_lines:
        para(line)
    if not intelligence_lines:
        para("No intelligence execution coverage was recorded.")
    para("This assessment applies only to the specified time range and the telemetry listed above. Absence of findings does not prove that the environment is free of malicious activity.")

    heading("MITRE ATT&CK Context")
    if attack_rows:
        attack=doc.add_table(rows=1,cols=3); attack.style="Light Shading Accent 1"; attack.alignment=WD_TABLE_ALIGNMENT.CENTER
        for index,label in enumerate(("Technique","Name","Tactic")):
            set_cell_text(attack.rows[0].cells[index],label,True,white); shade(attack.rows[0].cells[index],navy)
        mark_header(attack.rows[0])
        for technique_id,name,tactic in attack_rows:
            cells=attack.add_row().cells
            for index,value in enumerate((technique_id,name,tactic)): set_cell_text(cells[index],value,index==0)
        format_table(attack)
        para(f"Technique names are mapped to MITRE ATT&CK {ATTACK_RELEASE} ({ATTACK_RELEASE_DATE}).")
    else:
        para("No supported finding was mapped to a MITRE ATT&CK technique.")

    assessment_heading=heading("Assessment and Confidence"); assessment_heading.paragraph_format.page_break_before=True
    para(f"Overall assessment: {verdict}. Overall confidence: {overall_confidence.title()}.")
    confidence_bases=[str(item.get("confidence_justification") or "").strip() for item in findings]
    confidence_bases=[item for item in confidence_bases if item]
    if confidence_bases: para("Confidence basis: "+" ".join(confidence_bases))
    para("ROMULUS does not automatically declare activity confirmed malicious. Analyst validation is required before containment or remediation.")

    heading("Recommended Actions")
    for index,action in enumerate(actions,1): para(f"{index}. {action}")

    divider=doc.add_paragraph(); divider.paragraph_format.space_before=Pt(12)
    run=divider.add_run("Technical record"); run.bold=True; run.font.color.rgb=RGBColor.from_string(blue)
    divider.add_run("  The corresponding TXT report contains search history, hypotheses, tool activity, evidence references, source errors, and the complete candidate working list.")
    doc.save(path)
