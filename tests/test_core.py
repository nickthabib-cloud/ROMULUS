import importlib, json, os, sys, tempfile, types, unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch
from romulus.config import Settings
from romulus.intelligence.knowledge import hypotheses_for
from romulus.intelligence.correlate import correlate
from romulus.models import Evidence, encode
from romulus.store import Store
from romulus.tools.elasticsearch_client import validate_query
from romulus.tools.tshark_client import ALLOWED_FIELDS
from romulus.config import safe_url


@contextmanager
def _ollama_module():
    """Load the local-only adapter without requiring requests in the lean test runtime."""
    requests_stub=types.ModuleType("requests")
    requests_stub.Session=lambda: None
    sys.modules.pop("romulus.agent.ollama",None)
    with patch.dict(sys.modules,{"requests":requests_stub}):
        try:
            yield importlib.import_module("romulus.agent.ollama")
        finally:
            sys.modules.pop("romulus.agent.ollama",None)

class CoreTests(unittest.TestCase):
    def test_hypotheses_are_mission_relevant(self):
        hs=hypotheses_for("Hunt lateral movement over SMB")
        self.assertTrue(any(h.technique_id=="T1021.002" for h in hs))
    def test_time_window_filter_uses_explicit_scope_or_bounded_default(self):
        from romulus.investigation.workflow import _time_filter, _pcap_time_filter
        self.assertEqual(_time_filter("from 2026-01-01 to 2026-01-02",24),
            {"gte":"2026-01-01T00:00:00.000Z","lte":"2026-01-02T23:59:59.999Z"})
        self.assertEqual(_time_filter("from 2026-01-01 to 2026-01-02",24,"America/New_York"),
            {"gte":"2026-01-01T05:00:00.000Z","lte":"2026-01-03T04:59:59.999Z"})
        self.assertEqual(_time_filter("look at recent logs",12),{"gte":"now-12h","lte":"now"})
        self.assertIn("frame.time_epoch >=",_pcap_time_filter(["2024-11-13","2024-11-14"]))
    def test_host_query_uses_hypothesis_indicators_and_can_search_all_history(self):
        from romulus.investigation.workflow import _host_query
        hypotheses=[{"technique_id":"T1021.002"}]
        query=_host_query(hypotheses,["message","event.code","@timestamp"],"@timestamp",None,"Hunt for SMB lateral movement")
        encoded=str(query)
        self.assertIn("smb",encoded.lower()); self.assertIn("7045",encoded); self.assertNotIn("Hunt for SMB lateral movement",encoded)
        self.assertNotIn("range",query["bool"])
    def test_default_configuration_uses_production_safety_backstops(self):
        names=("MAX_RUNTIME_SECONDS","MAX_DEPTH","MAX_TOOL_CALLS","MAX_STAGNANT_CYCLES")
        with patch.dict(os.environ,{name:"" for name in names}):
            for name in names: os.environ.pop(name,None)
            cfg=Settings.from_env()
        self.assertTrue(cfg.historical_search)
        self.assertEqual((cfg.max_runtime_seconds,cfg.max_depth,cfg.max_tool_calls,
            cfg.max_stagnant_cycles),(3600,40,500,3))
        self.assertEqual(cfg.retention_policy,"compact")
        example=(Path(__file__).resolve().parents[1]/".env.example").read_text(encoding="utf-8")
        for setting in ("MAX_RUNTIME_SECONDS=3600","MAX_DEPTH=40","MAX_TOOL_CALLS=500",
                "MAX_STAGNANT_CYCLES=3"):
            self.assertIn(setting,example)
    def test_safety_backstops_clamp_to_documented_finite_ranges(self):
        names=("MAX_RUNTIME_SECONDS","MAX_DEPTH","MAX_TOOL_CALLS","MAX_STAGNANT_CYCLES")
        with patch.dict(os.environ,{name:"999999" for name in names}):
            high=Settings.from_env()
        self.assertEqual((high.max_runtime_seconds,high.max_depth,high.max_tool_calls,
            high.max_stagnant_cycles),(86400,200,5000,10))
        with patch.dict(os.environ,{name:"-999999" for name in names}):
            low=Settings.from_env()
        self.assertEqual((low.max_runtime_seconds,low.max_depth,low.max_tool_calls,
            low.max_stagnant_cycles),(10,1,1,1))
    def test_graph_recursion_limit_scales_with_supported_hunt_budgets(self):
        from dataclasses import replace
        from romulus.investigation.workflow import HuntEngine
        cfg=replace(Settings.from_env(),max_tool_calls=500,max_depth=40)
        engine=HuntEngine(cfg,Store(cfg.data_dir,cfg.logs_dir),offline=True)
        self.assertGreater(engine._graph_recursion_limit(),100)
        self.assertGreaterEqual(engine._graph_recursion_limit(),
            16+(3*cfg.max_tool_calls)+(3*cfg.max_depth))
    def test_scaled_limit_allows_real_graph_to_cross_100_steps(self):
        try:
            from langgraph.graph import StateGraph, END
        except ImportError:
            self.skipTest("Install requirements.txt for the real LangGraph recursion check")
        from dataclasses import replace
        from typing import TypedDict
        from romulus.investigation.workflow import HuntEngine
        class CounterState(TypedDict):
            count: int
        cfg=replace(Settings.from_env(),max_tool_calls=500,max_depth=40)
        engine=HuntEngine(cfg,Store(cfg.data_dir,cfg.logs_dir),offline=True)
        graph=StateGraph(CounterState)
        graph.add_node("step",lambda state:{"count":state["count"]+1})
        graph.set_entry_point("step")
        graph.add_conditional_edges("step",
            lambda state:"again" if state["count"]<110 else "done",
            {"again":"step","done":END})

        result=graph.compile().invoke({"count":0},
            config={"recursion_limit":engine._graph_recursion_limit()})

        self.assertEqual(result["count"],110)
    def test_url_redaction(self):
        self.assertEqual(safe_url("http://alice:secret@example.test:8005/api?token=hide"),"http://example.test:8005/api")
    def test_arkime_adapter_uses_digest_and_bounded_documented_query(self):
        try:
            from requests.auth import HTTPDigestAuth
            from romulus.tools.arkime_client import ArkimeAdapter
        except ImportError: self.skipTest("Install requirements.txt for Arkime adapter verification")
        from dataclasses import replace
        class Response:
            ok=True; status_code=200
            def json(self): return {"data":[{"id":"abc123"}]}
            def raise_for_status(self): pass
            def iter_content(self,chunk_size): return iter([b"pcap-bytes"])
        class Session:
            def __init__(self): self.calls=[]
            def get(self,*args,**kwargs): self.calls.append((args,kwargs)); return Response()
        sess=Session(); cfg=replace(Settings.from_env(),arkime_username="agent",arkime_password="secret",arkime_max_results=25,arkime_default_hours=24,historical_search=False)
        api=ArkimeAdapter(cfg,sess); found=api.search_sessions("port == 445",length=1000)
        self.assertEqual(len(found["sessions"]),1); self.assertEqual(sess.calls[-1][1]["params"]["length"],25)
        self.assertEqual(sess.calls[-1][1]["params"]["date"],24); self.assertIsInstance(sess.calls[-1][1]["auth"],HTTPDigestAuth)
        self.assertEqual(api.retrieve_pcap(["abc123"]),b"pcap-bytes")
        with self.assertRaises(ValueError): api.retrieve_pcap(["../../all"])
        api.check()
        self.assertTrue(sess.calls[-1][0][0].endswith("/api/sessions"))
        self.assertEqual(sess.calls[-1][1]["params"],{"date":-1,"length":1})
    def test_elasticsearch_adapter_is_allowlisted_and_caps_hits(self):
        from dataclasses import replace
        from romulus.tools.elasticsearch_client import ElasticsearchAdapter
        class Indices:
            def get(self,**kwargs): return {"ainception-hostlogs-a-b":{}}
            def get_mapping(self,**kwargs): return {"ainception-hostlogs-a-b":{"mappings":{"properties":{"@timestamp":{"type":"date"}}}}}
        class Client:
            indices=Indices()
            def search(self,**kwargs): return {"hits":{"hits":[{"_index":kwargs["index"],"_id":"1","_source":{"message":"synthetic"}}],"total":{"value":1}}}
        cfg=replace(Settings.from_env(),max_results=5)
        api=ElasticsearchAdapter(cfg,Client())
        self.assertTrue(api.indices()); self.assertIn("ainception-hostlogs-a-b",api.mappings("ainception-hostlogs-a-b"))
        self.assertTrue(api.check()["authenticated"])
        self.assertEqual(len(api.search("ainception-hostlogs-a-b",{"match":{"message":"synthetic"}},100)["hits"]),1)
        self.assertEqual(len(api.search_many(["ainception-hostlogs-a-b"],{"match":{"message":"synthetic"}},100)["hits"]),1)
        with self.assertRaises(ValueError): api.search("unapproved-index",{"match":{"message":"x"}},10)
    def test_tshark_wrapper_uses_allowlist_and_no_shell(self):
        from romulus.investigation.search_planner import SOURCE_FIELDS
        from romulus.tools.tshark_client import TSharkAdapter
        self.assertLessEqual(SOURCE_FIELDS["tshark"],ALLOWED_FIELDS)
        adapter=TSharkAdapter(Settings.from_env())
        class Process:
            stdout=["1\t10.0.0.1\n"]
            stderr=[]
            def wait(self,timeout): return 0
            def kill(self): pass
        with patch.object(adapter,"_approved",return_value=Path("approved.pcap")), patch("romulus.tools.tshark_client.shutil.which",return_value="tshark"), patch("romulus.tools.tshark_client.subprocess.Popen",return_value=Process()) as popen:
            out=adapter.inspect("approved.pcap",display_filter="tcp.port == 445",fields=["frame.number","ip.src"])
            self.assertEqual(out["packets"][0]["ip.src"],"10.0.0.1"); self.assertNotIn("shell",popen.call_args.kwargs)
            with self.assertRaises(ValueError): adapter.inspect("approved.pcap",display_filter="tcp.port == 445; rm -rf")
            with self.assertRaises(ValueError): adapter.inspect("approved.pcap",fields=["frame.payload"])
    def test_ollama_adapter_uses_local_api_and_caps_telemetry(self):
        try: from romulus.agent.ollama import Ollama
        except ImportError: self.skipTest("Install requirements.txt for Ollama adapter verification")
        class Response:
            def __init__(self,body): self.body=body
            def raise_for_status(self): pass
            def json(self): return self.body
        class Session:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs): self.payload=kwargs["json"]; return Response({"message":{"content":"Evidence remains inconclusive."}})
        session=Session(); model=Ollama(Settings.from_env(),session)
        note=model.summarize("Hunt SMB",["synthetic data"]*25)
        self.assertIn("inconclusive",note); self.assertEqual(len(session.payload["messages"][1]["content"].split("UNTRUSTED TELEMETRY DATA:")),2)
        self.assertEqual(session.payload["options"]["num_ctx"],6144)
        class PlanningSession:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs):
                self.payload=kwargs["json"]
                content='{"searches":[{"source":"elasticsearch","hypothesis_id":"H1","intent":"Find PowerShell execution","terms":["powershell"],"fields":["process.name"],"ips":[],"hostnames":[],"ports":[],"protocols":[],"event_codes":[],"index_class":"host","anchor_evidence_ids":[],"native_query_language":"none","native_query":""}]}'
                return Response({"message":{"content":content}})
        planning_session=PlanningSession()
        searches=Ollama(Settings.from_env(),planning_session).propose_searches(
            "Hunt PowerShell",[{"hypothesis_id":"H1","title":"PowerShell execution"}],[],
            {"elasticsearch":["process.name"]},1)
        self.assertEqual(searches[0]["intent"],"Find PowerShell execution")
        schema=planning_session.payload["format"]
        self.assertEqual(schema["properties"]["searches"]["items"]["properties"]["source"]["enum"],["elasticsearch"])
        self.assertIn("candidate_id",schema["properties"]["searches"]["items"]["properties"])
        self.assertIn("follow_up_id",schema["properties"]["searches"]["items"]["properties"])
        self.assertEqual(schema["required"],["searches"]); self.assertFalse(schema["additionalProperties"])
        self.assertEqual(planning_session.payload["options"]["temperature"],0.15)
        class ReasoningSession:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs):
                self.payload=kwargs["json"]
                content='{"assessments":[],"findings":[],"new_hypotheses":[],"searches":[],"continue_investigation":false,"reasoning":"No supported expansion."}'
                return Response({"message":{"content":content}})
        reasoning_session=ReasoningSession()
        result=Ollama(Settings.from_env(),reasoning_session).reason_hunt(
            "Hunt for threats",[],[],{
                "search_fields":{"elasticsearch":["process.name"]},
                "model_initial_search_plan":"deferred: undirected mission has no telemetry-backed candidate",
                "candidate_queue":[{"candidate_id":"CL-1","evidence_ids":["EV-1"]}],
            },[],1)
        self.assertFalse(result["continue_investigation"])
        prompt=reasoning_session.payload["messages"][1]["content"]
        self.assertIn("deferred: undirected mission",prompt)
        self.assertIn("CL-1",prompt)
        reason_schema=reasoning_session.payload["format"]
        self.assertEqual(reason_schema["properties"]["new_hypotheses"]["maxItems"],0)
        finding_schema=reason_schema["properties"]["findings"]["items"]
        self.assertIn("hypothesis_id",finding_schema["required"])
    def test_investigator_schema_allows_three_grounded_competing_hypotheses_with_falsifiers(self):
        with _ollama_module() as module:
            schema=module._investigator_schema(["elasticsearch","arkime"],["H-EXISTING"],3,3)
        self.assertEqual(set(schema["properties"]),{
            "assessments","new_hypotheses","searches","findings","summary"})
        self.assertEqual(schema["properties"]["new_hypotheses"]["maxItems"],3)
        hypothesis=schema["properties"]["new_hypotheses"]["items"]
        for required in ("candidate_id","evidence_ids","title","technique_id",
                "explanation","expected_host","expected_network","falsifying_observations",
                "benign_explanations","analytic_goal"):
            self.assertIn(required,hypothesis["required"])
        self.assertEqual(hypothesis["properties"]["falsifying_observations"]["minItems"],1)
        self.assertEqual(hypothesis["properties"]["benign_explanations"]["minItems"],1)
        self.assertIn("anyOf",hypothesis)
        for collection in ("assessments","new_hypotheses","searches","findings"):
            item=schema["properties"][collection]["items"]
            self.assertNotIn("proposal_id",item["required"])
            self.assertNotIn("proposal_id",item["properties"])
        search=schema["properties"]["searches"]["items"]
        self.assertIn("hypothesis_ref",search["required"])
        self.assertNotIn("hypothesis_id",search["properties"])

    def test_investigator_assigns_stable_unique_proposal_ids_after_parsing(self):
        structured={
            "assessments":[{"hypothesis_id":"H-1","status":"inconclusive","evidence_ids":["EV-1"],
                "observation":"One process record is visible.","interpretation":"The cause is unresolved.",
                "benign_alternatives":["Authorized administration"],"rationale":"More context is required."}],
            "new_hypotheses":[{"candidate_id":"CL-1","evidence_ids":["EV-1"],
                "title":"Unexpected process lineage near the retained event","technique_id":"T1059.001",
                "explanation":"The cited process may have an unexpected parent.",
                "expected_host":["An unusual parent process"],"expected_network":[],
                "falsifying_observations":["The lineage belongs to an approved job"],
                "benign_explanations":["Authorized administration"],
                "analytic_goal":"Determine whether the process lineage is expected."}],
            "searches":[{"source":"elasticsearch","hypothesis_ref":"H-1",
                "intent":"Inspect the cited process lineage.","direction":"backward",
                "discriminating_goal":"Separate approved administration from unexpected execution.",
                "candidate_id":"CL-1","follow_up_id":"FU-1","terms":["powershell.exe"],
                "fields":["process.name"],"ips":[],"hostnames":[],"ports":[],"protocols":[],
                "event_codes":[],"index_class":"host","anchor_evidence_ids":["EV-1"],
                "native_query_language":"none","native_query":""},
                {"source":"elasticsearch","hypothesis_ref":"NEW-HYPOTHESIS-1",
                "intent":"Test the proposed process-lineage explanation.","direction":"validation",
                "discriminating_goal":"Distinguish an approved job from unexpected execution.",
                "candidate_id":"CL-1","follow_up_id":"FU-1","terms":["powershell.exe"],
                "fields":["process.name"],"ips":[],"hostnames":[],"ports":[],"protocols":[],
                "event_codes":[],"index_class":"host","anchor_evidence_ids":["EV-1"],
                "native_query_language":"none","native_query":""}],
            "findings":[{"hypothesis_id":"H-1","title":"Process lineage requires review",
                "technique_id":"T1059.001","observation":"The cited event has unusual lineage context.",
                "interpretation":"The activity remains a suspected lead.","evidence_ids":["EV-1"],
                "confidence":"low","confidence_justification":"Only one record is visible.",
                "alternatives":["Authorized administration"],"follow_up":["Validate the parent process."]}],
            "summary":"One bounded follow-up is warranted.",
        }
        class Response:
            def raise_for_status(self): pass
            def json(self): return {"message":{"content":json.dumps(structured)}}
        class Session:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs): self.payload=kwargs["json"]; return Response()
        case_file={"version":1,"candidates":[{"candidate_id":"CL-1","evidence_ids":["EV-1"],
                "assets":["workstation-1"],"attack_techniques":["T1059.001"],
                "follow_ups":[{"follow_up_id":"FU-1","status":"pending"}]}],
            "hypotheses":[{"hypothesis_id":"H-1","candidate_id":"CL-1","evidence_ids":["EV-1"],
                "technique_id":"T1059.001","status":"inconclusive"}],
            "evidence":[{"evidence_id":"EV-1","asset":"workstation-1",
                "observed_at":"2026-01-01T00:00:00Z",
                "observation":"A bounded synthetic process event.",
                "raw":{"process.name":"powershell.exe"}}]}

        with _ollama_module() as module:
            first_session=Session(); second_session=Session()
            first=module.Ollama(Settings.from_env(),first_session).investigate_case(
                case_file,{"elasticsearch":["process.name"]},3,3)
            second=module.Ollama(Settings.from_env(),second_session).investigate_case(
                case_file,{"elasticsearch":["process.name"]},3,3)

        first_ids=[row["proposal_id"] for key in
            ("assessments","new_hypotheses","searches","findings") for row in first[key]]
        second_ids=[row["proposal_id"] for key in
            ("assessments","new_hypotheses","searches","findings") for row in second[key]]
        self.assertEqual(first_ids,second_ids)
        self.assertEqual(len(first_ids),len(set(first_ids)))
        self.assertTrue(all(identifier.startswith("P-") for identifier in first_ids))
        self.assertEqual(first["searches"][1]["hypothesis_ref"],
            first["new_hypotheses"][0]["proposal_id"])
        search_ref_values=first_session.payload["format"]["properties"]["searches"]["items"] \
            ["properties"]["hypothesis_ref"]["enum"]
        self.assertIn("H-1",search_ref_values)
        self.assertIn("NEW-HYPOTHESIS-1",search_ref_values)
        critic={"reviews":[{"proposal_id":proposal_id,"decision":"accept",
            "issue_codes":[],"rationale":"The bounded proposal is grounded."}
            for proposal_id in first_ids]}
        from romulus.investigation.reasoning import admit_reasoning
        admission=admit_reasoning(case_file,first,critic,case_file["hypotheses"])
        self.assertEqual(len(admission["new_hypotheses"]),1)
        admitted_hypothesis_id=admission["new_hypotheses"][0]["hypothesis_id"]
        self.assertTrue(any(search["hypothesis_id"]==admitted_hypothesis_id
            for search in admission["searches"]))
        self.assertNotIn("proposal_id",first_session.payload["format"]["properties"]
            ["assessments"]["items"]["properties"])
        self.assertGreaterEqual(first_session.payload["options"]["num_ctx"],16384)
        self.assertGreaterEqual(first_session.payload["options"]["num_predict"],4000)

    def test_structured_adapter_preserves_complete_json_beyond_legacy_slice(self):
        large_value="x"*26000
        class Response:
            def raise_for_status(self): pass
            def json(self): return {"done_reason":"stop",
                "message":{"content":json.dumps({"payload":large_value})}}
        class Session:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs): self.payload=kwargs["json"]; return Response()
        schema={"type":"object","properties":{"payload":{"type":"string"}},
            "required":["payload"],"additionalProperties":False}

        with _ollama_module() as module:
            result=module.Ollama(Settings.from_env(),Session())._json_chat(
                "system","prompt",schema,tokens=5000)

        self.assertEqual(result["payload"],large_value)

    def test_structured_adapter_still_rejects_provider_truncation_and_incomplete_json(self):
        class Response:
            def __init__(self,body): self.body=body
            def raise_for_status(self): pass
            def json(self): return self.body
        class Session:
            def __init__(self,body): self.body=body
            def post(self,url,**kwargs): return Response(self.body)
        schema={"type":"object","properties":{"value":{"type":"string"}},
            "required":["value"],"additionalProperties":False}

        with _ollama_module() as module:
            with self.assertRaisesRegex(ValueError,"truncated"):
                module.Ollama(Settings.from_env(),Session({"done_reason":"length",
                    "message":{"content":"{\"value\":\"complete\"}"}}))._json_chat(
                        "system","prompt",schema,tokens=100)
            with self.assertRaisesRegex(ValueError,"invalid, truncated, or incomplete"):
                module.Ollama(Settings.from_env(),Session({"done_reason":"stop",
                    "message":{"content":"{\"value\":\"cut"}}))._json_chat(
                        "system","prompt",schema,tokens=100)
    def test_investigator_prompt_receives_only_the_case_file_and_marks_telemetry_untrusted(self):
        class Response:
            def raise_for_status(self): pass
            def json(self):
                return {"message":{"content":('{"assessments":[],"new_hypotheses":[],"searches":[],'
                    '"findings":[],"summary":"No admissible proposal."}')}}
        class Session:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs): self.payload=kwargs["json"]; return Response()
        case_file={"version":1,"mission":{"hunt_id":"HUNT-CASE"},
            "evidence":[{"evidence_id":"EV-VISIBLE","observation":"CASE-FILE-SENTINEL"}]}
        session=Session()

        with _ollama_module() as module:
            result=module.Ollama(Settings.from_env(),session).investigate_case(
                case_file,{"elasticsearch":["process.name"]},max_searches=3,max_hypotheses=3)

        self.assertEqual(result["summary"],"No admissible proposal.")
        system=session.payload["messages"][0]["content"].lower()
        prompt=session.payload["messages"][1]["content"]
        self.assertIn("untrusted",system)
        self.assertIn("untrusted",prompt.lower())
        self.assertIn("CASE-FILE-SENTINEL",prompt)
        self.assertIn('"EV-VISIBLE"',prompt)
        self.assertNotIn("credentials",prompt.lower())
        self.assertLessEqual(session.payload["options"]["temperature"],0.15)
    def test_critic_schema_can_only_review_known_proposal_ids(self):
        investigator={
            "assessments":[{"proposal_id":"P-ASSESS"}],
            "new_hypotheses":[{"proposal_id":"P-HYP"}],
            "searches":[{"proposal_id":"P-SEARCH"}],
            "findings":[{"proposal_id":"P-FIND"}],
        }

        with _ollama_module() as module:
            schema=module._critic_schema(investigator)

        self.assertEqual(set(schema["properties"]),{"reviews"})
        review=schema["properties"]["reviews"]["items"]
        self.assertEqual(set(review["properties"]),{"proposal_id","decision","issue_codes","rationale"})
        self.assertEqual(set(review["properties"]["proposal_id"]["enum"]),
            {"P-ASSESS","P-HYP","P-SEARCH","P-FIND"})
        self.assertEqual(review["properties"]["decision"]["enum"],["accept","reject"])
        self.assertEqual(set(review["properties"]["issue_codes"]["items"]["enum"]),{
            "citation_mismatch","unsupported_causality","ignored_benign_alternative",
            "duplicate_work","missing_discriminating_search","avoidable_search_cost","malformed"})
        self.assertFalse(schema["additionalProperties"])
        self.assertFalse(review["additionalProperties"])
    def test_critic_prompt_cannot_author_evidence_searches_findings_or_priorities(self):
        class Response:
            def raise_for_status(self): pass
            def json(self):
                return {"message":{"content":('{"reviews":[{"proposal_id":"P-1",'
                    '"decision":"accept","issue_codes":[],"rationale":"Grounded."}]}')}}
        class Session:
            def __init__(self): self.payload=None
            def post(self,url,**kwargs): self.payload=kwargs["json"]; return Response()
        investigator={"assessments":[{"proposal_id":"P-1","hypothesis_id":"H-1"}],
            "new_hypotheses":[],"searches":[],"findings":[],"summary":"bounded"}
        session=Session()

        with _ollama_module() as module:
            result=module.Ollama(Settings.from_env(),session).criticize_case(
                {"version":1,"evidence":[{"evidence_id":"EV-1"}]},investigator)

        self.assertEqual(result["reviews"][0]["proposal_id"],"P-1")
        prompt=session.payload["messages"][1]["content"].lower()
        for phrase in ("cannot author new evidence","cannot author new hypotheses",
                "cannot author new searches","cannot author new findings","cannot assign priority"):
            self.assertIn(phrase,prompt)
        self.assertEqual(set(session.payload["format"]["properties"]),{"reviews"})
    def test_reasoning_preflight_validates_both_structured_contracts(self):
        with _ollama_module() as module:
            result=module.Ollama(Settings.from_env()).reasoning_preflight()

        self.assertTrue(result["ready"])
        self.assertTrue(result["investigator_ready"])
        self.assertTrue(result["critic_ready"])
        self.assertEqual(result["investigator_schema_keys"],
            ["assessments","findings","new_hypotheses","searches","summary"])
        self.assertEqual(result["critic_schema_keys"],["reviews"])
    def test_elasticsearch_query_validation(self):
        validate_query({"bool":{"should":[{"match":{"message":"SMB"}}],"minimum_should_match":1}})
        with self.assertRaises(ValueError): validate_query({"query_string":{"query":"*"}})
    def test_correlation_needs_shared_ip_and_time(self):
        h={"evidence_id":"h","observed_at":"2026-01-01T00:00:00Z","raw":{"ip":"10.0.0.2","event_id":4624}}
        n={"evidence_id":"n","observed_at":"2026-01-01T00:01:00Z","raw":{"src_ip":"10.0.0.2","dst_port":445,"protocol":"smb"}}
        self.assertEqual(len(correlate([h],[n])),1)
        n["raw"]["src_ip"]="10.0.0.3"; self.assertEqual(correlate([h],[n]),[])
        n["raw"]["src_ip"]="10.0.0.2"; h["raw"].pop("event_id"); self.assertEqual(correlate([h],[n]),[])
    def test_store_roundtrip_and_event_log(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as d:
            p=Path(d); st=Store(p/"data",p/"logs"); st.save_state("H1",{"status":"partial"})
            self.assertEqual(st.load_state("H1")["status"],"partial")
            st.upsert_finding("H1",{"finding_id":"F1","status":"suspected"})
            self.assertIn('"event": "created"',st.events_path.read_text())
            st.audit({"hunt_id":"H1","parameters":{"nested":{"api_token":"do-not-write"}},"result_count":0})
            self.assertNotIn("do-not-write",st.audit_path.read_text()); self.assertEqual(st.audit_records("H1")[0]["result_count"],0)
    def test_reasoning_depth_limit_wins_over_minimum_rounds(self):
        from dataclasses import replace
        from romulus.investigation.workflow import HuntEngine
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as d:
            p=Path(d); cfg=replace(Settings.from_env(),home=p,reports_dir=p/"reports",data_dir=p/"data",logs_dir=p/"logs",
                max_depth=2,min_reasoning_rounds=3)
            state={"hunt_id":"H-DEPTH","objective":"Hunt for SMB lateral movement",
                "mission_mode":"directed","depth":2,"tool_calls":0,"started_epoch":__import__("time").time(),
                "coverage":{},"search_tasks":{},"search_queue":[],"search_outcomes":[],
                "scheduler_history":[],"candidate_leads":[]}
            engine=HuntEngine(cfg,Store(cfg.data_dir,cfg.logs_dir),offline=True)
            self.assertEqual(engine._after_reason(state),"report")
            self.assertEqual(state["termination_reason"],"depth_budget")
    def test_empty_partial_reports_are_generated(self):
        try: import docx
        except ImportError: self.skipTest("Install requirements.txt for Word report verification")
        from romulus.reporting.report import generate_report
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as d:
            out=Path(d); paths=generate_report({"hunt_id":"HUNT-EMPTY","status":"partial","objective":"empty synthetic test\x01","summary":"No data was returned.\x01","coverage":{"host":"unavailable\x01"},"host":[{"evidence_id":"E1","observation":"malformed\x0b telemetry"}]},out)
            self.assertEqual({p.suffix for p in paths},{".docx",".txt"})
            self.assertIn("does not establish",paths[1].read_text()); self.assertNotIn("\x01",paths[1].read_text()); self.assertNotIn("\x0b",paths[1].read_text())
            from docx import Document
            doc=Document(next(p for p in paths if p.suffix==".docx"))
            body="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            self.assertNotIn("\x01",body); self.assertNotIn("\x0b",body)
            self.assertNotIn("malformed telemetry",body)
            self.assertIn("malformed telemetry",paths[1].read_text())
    def test_synthetic_full_investigation_and_reports(self):
        try:
            import langgraph,docx
        except ImportError: self.skipTest("Install requirements.txt for the full offline graph and Word report")
        from romulus.investigation.workflow import HuntEngine
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as d:
            p=Path(d); base=Settings.from_env()
            from dataclasses import replace
            cfg=replace(base,home=p,reports_dir=p/"reports",data_dir=p/"data",logs_dir=p/"logs")
            store=Store(cfg.data_dir,cfg.logs_dir); state=HuntEngine(cfg,store,offline=True).run("Hunt for lateral movement using SMB")
            self.assertEqual(state["status"],"completed"); self.assertTrue(state["findings"])
            self.assertEqual(state["outcome"],"evidence_supported_lead")
            self.assertEqual(state["termination_reason"],"coverage_closure_reached")
            self.assertEqual(state["retention"]["status"],"completed")
            self.assertNotIn("host",state); self.assertNotIn("network",state)
            cited={evidence_id for finding in state["findings"] for evidence_id in finding.get("evidence_ids",[])}
            self.assertEqual({item["evidence_id"] for item in store.list_evidence(state["hunt_id"])},cited)
            import sqlite3
            db=sqlite3.connect(store.checkpoint_path)
            try: self.assertEqual(db.execute("SELECT COUNT(*) FROM checkpoints WHERE thread_id=?",(state["hunt_id"],)).fetchone()[0],0)
            finally: db.close()
            self.assertEqual(store.audit_records(state["hunt_id"]),[])
            events=(cfg.logs_dir/"findings.jsonl").read_text() if (cfg.logs_dir/"findings.jsonl").exists() else ""
            self.assertNotIn(state["hunt_id"],events)
            self.assertTrue(any(p.suffix==".docx" for p in map(Path,state["report_paths"])))
            self.assertTrue(any(p.suffix==".txt" for p in map(Path,state["report_paths"])))
            from docx import Document
            doc=Document(next(Path(p) for p in state["report_paths"] if p.endswith(".docx")))
            body="\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Executive Summary",body); self.assertIn("Correlated SMB",body); self.assertIn("Supported Attack Path",body)

    def test_live_workflow_runs_bounded_historical_pivots_and_audits_every_call(self):
        try: import langgraph,docx
        except ImportError: self.skipTest("Install requirements.txt for the LangGraph integration check")
        from dataclasses import replace
        from datetime import datetime,timezone
        from romulus.investigation.workflow import HuntEngine
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as d:
            p=Path(d); cfg=replace(Settings.from_env(),home=p,reports_dir=p/"reports",data_dir=p/"data",logs_dir=p/"logs",
                arkime_username="agent",arkime_password="secret",pcap_files=(),pcap_allowed_dirs=(),
                max_tool_calls=40,max_depth=4,min_reasoning_rounds=1,historical_search=True)
            class ES:
                def __init__(self): self.searches=0
                def indices(self): return [{"index":"ainception-hostlogs-a-b"}]
                def mappings_many(self,indices): return {"ainception-hostlogs-a-b":{"mappings":{"properties":{
                    "@timestamp":{"type":"date"},"message":{"type":"text"},"event":{"properties":{"code":{"type":"long"}}},
                    "source":{"properties":{"ip":{"type":"ip"}}},"host":{"properties":{"name":{"type":"keyword"}}}}}}}
                def search_many(self,indices,query,size,**kwargs):
                    self.searches+=1
                    code=4624 if self.searches==1 else 7045
                    msg="successful 4624 logon" if code==4624 else "7045 service installation"
                    return {"hits":[{"_index":indices[0],"_id":str(self.searches),"@timestamp":"2024-11-13T10:00:00Z",
                        "message":msg,"event":{"code":code},"source":{"ip":"10.0.0.5"},"host":{"name":"ws-01"}}]}
            class Arkime:
                def __init__(self): self.calls=[]
                def search_sessions_sampled(self,expression,**kwargs):
                    self.calls.append((expression,kwargs))
                    stamp=datetime(2024,11,13,10,0,tzinfo=timezone.utc).timestamp()
                    return {"sessions":[{"id":f"session-{len(self.calls)}","source.ip":"10.0.0.5","destination.ip":"10.0.0.9",
                        "destination.port":445,"firstPacket":stamp}],"total":1,"capped":False,
                        "sampling":{"version":1,"strategy":"exhaustive","requested_strata":1,
                            "covered_strata":1,"retained_by_stratum":[1]},"provenance":{"expression":expression}}
            store=Store(cfg.data_dir,cfg.logs_dir); engine=HuntEngine(cfg,store,use_ollama=False)
            es=ES(); arkime=Arkime(); engine.es=es; engine.arkime=arkime
            state=engine.run("Hunt for SMB lateral movement")
            self.assertEqual(state["status"],"partial")
            self.assertFalse(state["coverage"]["hypothesis_coverage"]["closure_ready"],
                state["coverage"]["hypothesis_coverage"])
            self.assertEqual(state["coverage"]["time_scope"],"all retained indexed time")
            self.assertEqual(state["source_health"],{"host":True,"network":True})
            executed_task_ids={outcome["task_id"] for outcome in state["search_outcomes"]}
            self.assertTrue(state["completed_leads"])
            self.assertTrue(set(state["completed_leads"])<=executed_task_ids)
            self.assertEqual({"host","network"},{outcome["source_family"] for outcome in state["search_outcomes"]})
            from romulus.investigation.scheduler import outcome_resolves
            self.assertFalse(any(outcome_resolves(outcome) for outcome in state["search_outcomes"]
                if outcome["source_family"]=="host"))
            self.assertTrue(any(outcome_resolves(outcome) for outcome in state["search_outcomes"]
                if outcome["source_family"]=="network"))
            self.assertGreaterEqual(len(state["host"]),2); self.assertGreaterEqual(len(state["network"]),2)
            self.assertTrue(all(call[1].get("all_history") for call in arkime.calls))
            self.assertLessEqual(state["tool_calls"],cfg.max_tool_calls)
            self.assertEqual(state["tool_calls"],len(store.audit_records(state["hunt_id"])))
            self.assertTrue(all("result_count" in item and "time_range" in item and "duration_ms" in item for item in state["tool_runs"]))
            self.assertEqual(state["termination_reason"],"no_productive_frontier")

if __name__=="__main__": unittest.main()
